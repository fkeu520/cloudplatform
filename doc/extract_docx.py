import docx

doc = docx.Document(r'D:\work\AI\output\platform\doc\产品需求说明书.DOCX')

lines = []
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if text:
        style = para.style.name
        if 'Title' in style:
            lines.append(f'# {text}')
        elif 'Heading' in style:
            level = style.replace('Heading ', '')
            lines.append(f'{"#" * int(level)} {text}')
        elif 'List' in style:
            lines.append(f'- {text}')
        else:
            lines.append(text)
        lines.append('')

with open(r'D:\work\AI\output\platform\doc\产品需求说明书_full.md', 'w', encoding='utf-8') as f:
    f.write('\n'.join(lines))

# Extract tables
with open(r'D:\work\AI\output\platform\doc\产品需求说明书_tables.md', 'w', encoding='utf-8') as f:
    for ti, table in enumerate(doc.tables):
        f.write(f'\n## Table {ti+1}\n\n')
        for ri, row in enumerate(table.rows):
            cells = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
            f.write('| ' + ' | '.join(cells) + ' |\n')
        f.write('\n')

print(f'Total paragraphs: {len(doc.paragraphs)}')
print(f'Total tables: {len(doc.tables)}')
print('Done!')
